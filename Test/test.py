from openpyxl import Workbook, load_workbook
from datetime import datetime, timedelta
import re

def generate_time_intervals():
    start_time = datetime.strptime("00:00", "%H:%M")
    end_time = datetime.strptime("23:59", "%H:%M")
    delta = timedelta(minutes=15)

    times = []
    current = start_time
    while current <= end_time:
        times.append([current.strftime("%H:%M")])
        current += delta
    return times
time_list = generate_time_intervals()


day = []
for i in range(31):
    day.append(str(i+1))
month = []


# print(day)
# print(time_list)
# print(type(time_list[1][0]))
wb = load_workbook("record document/"+'13'+".xlsx")
ws = wb.active
for row in ws.iter_rows(min_row=4,min_col=3):
   for cell in row:
       # print(cell.column)
        continue
for row in ws.iter_rows(min_row=6,max_col=1):
    for cell in row:
        # print(cell.row)
        continue





current_time = datetime.now()
for i in ws['C4:AG4'][0]:
    if i.value == current_time.strftime("%d"):
        col = i.column
        print(col)
        print(type(col))
for i in ws['A6:A101']:
    if i[0].value > current_time.strftime("%H:%M"):
        row = i[0].row
        print(row)
        break


# for row in ws.values:
#    for value in row:
#      print(value)

# if '12:00' >current_time.strftime("%H:%M"):
#     print(True)
# for i in ws['A6:A101']:
#     print(i[0].value)
    # print(i.value > current_time.strftime("%H:%M"))