import logging
from openpyxl import load_workbook
from datetime import datetime


def setup_logger():
    logging.basicConfig(
        filename='excel_operations.log',
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s'
    )
    return logging.getLogger('ExcelLogger')


class ExcelFileHandler:
    def __init__(self, file_path):
        self.logger = setup_logger()
        self.file_path = file_path
        self.workbook = None

    def read_excel(self, sheet_name=None):
        try:
            start_time = datetime.now()
            self.workbook = load_workbook(self.file_path)
            sheet = self.workbook[sheet_name] if sheet_name else self.workbook.active

            self.logger.info(f"READ - File: {self.file_path}, Sheet: {sheet.title}, "
                             f"Duration: {(datetime.now() - start_time).total_seconds():.2f}s")
            return sheet
        except Exception as e:
            self.logger.error(f"READ FAILED - File: {self.file_path}, Error: {str(e)}")
            raise

    def write_excel(self, data, sheet_name="Sheet1"):
        try:
            start_time = datetime.now()
            if not self.workbook:
                self.workbook = load_workbook(self.file_path)

            if sheet_name in self.workbook.sheetnames:
                sheet = self.workbook[sheet_name]
            else:
                sheet = self.workbook.create_sheet(sheet_name)

            # Example write operation - modify as needed
            for row in data:
                sheet.append(row)

            self.workbook.save(self.file_path)
            self.logger.info(f"WRITE - File: {self.file_path}, Sheet: {sheet_name}, "
                             f"Rows: {len(data)}, Duration: {(datetime.now() - start_time).total_seconds():.2f}s")
        except Exception as e:
            self.logger.error(f"WRITE FAILED - File: {self.file_path}, Error: {str(e)}")
            raise
