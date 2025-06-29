from openpyxl import Workbook, load_workbook
from datetime import datetime
import re

def read_from_excel(filename):
    wb = load_workbook(filename)
    ws = wb.active

    data = []
    for row in ws.iter_rows(values_only=True):
        data.append(row)

    return data
def read_and_write(import_text):
    number = re.findall(r'\d+', import_text)[0]
    file_path = "record document/"+number+".xlsx"
    wb = load_workbook(file_path)
    ws = wb.active
    current_time = datetime.now()
    #Find the cell
    def Find_the_cell():
        global row, col
        for i in ws['C4:AG4'][0]:
            if i.value == current_time.strftime("%d"):
                col = i.column
                print(col)
                print(type(col))
        for i in ws['A6:A101']:
            if i[0].value > current_time.strftime("%H:%M"):
                row = i[0].row
                print(row)
                break
        return [row,col]
    #Write into the cell
    def write_to_cell(file_path, row, column, value):

        ws.cell(row=row, column=column, value=value)
        wb.save(file_path)

    def input_value():
        return '✓ '+current_time.strftime('%Y/%m/%d,%H:%M:%S')

    write_to_cell(file_path,*Find_the_cell(),input_value())


read_and_write('13號')

