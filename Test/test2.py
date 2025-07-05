from fastapi import FastAPI, HTTPException
from pydantic import BaseModel
import requests
import base64
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook, load_workbook
import re
import logging
import traceback

app = FastAPI(title="Cantonese Transcription API")


# data logger
class ExcelDataLogger:
    def __init__(self, log_file='excel_data_operations.log'):
        self.logger = logging.getLogger('ExcelDataLogger')
        self.logger.setLevel(logging.INFO)

        # Create file handler
        file_handler = logging.FileHandler(log_file)
        file_handler.setLevel(logging.INFO)

        # Create console handler for errors
        console_handler = logging.StreamHandler()
        console_handler.setLevel(logging.ERROR)

        # Create formatter
        formatter = logging.Formatter(
            '%(asctime)s - %(levelname)s - %(message)s\n'
        )
        file_handler.setFormatter(formatter)
        console_handler.setFormatter(formatter)

        # Add handlers
        self.logger.addHandler(file_handler)
        self.logger.addHandler(console_handler)

    def log_operation(self, operation, file_path, details=None, success=True, error=None):
        log_entry = {
            'operation': operation,
            'file': file_path,
            'timestamp': datetime.now().isoformat(),
            'success': success,
            'details': details or {},
            'error': str(error) if error else None
        }

        if success:
            self.logger.info(log_entry)
        else:
            self.logger.error(log_entry)
            if error:
                self.logger.error(traceback.format_exc())


class TranscriptionRequest(BaseModel):
    audio_data: str  # base64 encoded audio
    format: str  # audio format (mp3/wav)


class TranscriptionResponse(BaseModel):
    text: str


chinese_to_arabic = {
    '零': 0, '一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9,
    '十': 10, '百': 100, '千': 1000
}


def chinese_to_arabic_number(chinese_num):
    if not chinese_num:
        return 0

    total = 0
    current = 0
    prev_value = 0

    for char in chinese_num:
        value = chinese_to_arabic.get(char, 0)

        if value == 0:
            continue  # skip '零' (zero)
        elif value >= 10:  # multiplier (十, 百, 千, etc.)
            if current == 0:
                current = 1  # e.g., "十" means 10, not 0*10
            total += current * value
            current = 0
        else:
            current = value

    total += current
    return total


def extract_and_convert_chinese_numbers(text):
    # Regex to match Chinese numbers (e.g., "一百二十", "三")
    pattern = re.compile(r'[零一二三四五六七八九十百千万亿]+')
    matches = pattern.findall(text)

    converted_numbers = []
    for match in matches:
        arabic_num = chinese_to_arabic_number(match)
        converted_numbers.append(arabic_num)
        text = text.replace(match, str(arabic_num), 1)  # replace first occurrence

    return text, converted_numbers


def read_and_write(import_text):
    logger = ExcelDataLogger()

    try:
        # Extract number and create file path
        numbers = re.findall(r'\d+', import_text)
        if not numbers:
            raise ValueError("No numbers found in input text")

        number = numbers[0]
        file_path = "record document/" + number + ".xlsx"

        # Log file detection
        logger.log_operation(
            operation="file_detection",
            file_path=file_path,
            details={'extracted_number': number}
        )

        # Workbook operations
        try:
            start_time = datetime.now()
            wb = load_workbook(file_path)
            ws = wb.active
            current_time = datetime.now()

            logger.log_operation(
                operation="workbook_load",
                file_path=file_path,
                details={
                    'load_time': (datetime.now() - start_time).total_seconds(),
                    'sheet': ws.title
                }
            )

            def Find_the_cell():
                cell_find_start = datetime.now()
                current_day = current_time.strftime("%d")  # "03"
                current_day_no_zero = str(int(current_day))  # "3"
                current_time_str = current_time.strftime("%H:%M")

                # Search for day in header row (C4:AG4)
                col = None
                for cell in ws['C4:AG4'][0]:
                    cell_value = str(cell.value).strip() if cell.value else ""
                    # Match both "03" and "3" formats
                    if cell_value in [current_day, current_day_no_zero]:
                        col = cell.column
                        logger.log_operation(
                            operation="column_found",
                            file_path=file_path,
                            details={
                                'column': col,
                                'search_value': current_day,
                                'range_searched': 'C4:AG4'
                            }
                        )
                        break

                if col is None:
                    available_days = [str(cell.value).strip() for cell in ws['C4:AG4'][0] if cell.value]
                    raise ValueError(
                        f"Day '{current_day}' or '{current_day_no_zero}' not found in header row. "
                        f"Available days: {', '.join(available_days)}"
                    )

                # Search for time in first column (A6:A101)
                row = None
                for cell in ws['A6:A101']:
                    cell_value = cell[0].value
                    if cell_value and str(cell_value) > current_time_str:
                        row = cell[0].row
                        logger.log_operation(
                            operation="row_found",
                            file_path=file_path,
                            details={
                                'row': row,
                                'search_value': current_time_str,
                                'range_searched': 'A6:A101'
                            }
                        )
                        break

                if row is None:
                    available_times = [str(cell[0].value) for cell in ws['A6:A101'] if cell[0].value]
                    raise ValueError(
                        f"Time '{current_time_str}' not found in first column. "
                        f"Available times: {', '.join(available_times)}"
                    )

                logger.log_operation(
                    operation="cell_located",
                    file_path=file_path,
                    details={
                        'duration': (datetime.now() - cell_find_start).total_seconds(),
                        'coordinates': f"Row: {row}, Column: {col}"
                    }
                )
                return row, col

            def write_to_cell(file_path, row, column, value):
                write_start = datetime.now()
                try:
                    ws.cell(row=row, column=column, value=value)
                    wb.save(file_path)

                    logger.log_operation(
                        operation="cell_write",
                        file_path=file_path,
                        details={
                            'coordinates': f"Row: {row}, Column: {column}",
                            'value_written': value,
                            'duration': (datetime.now() - write_start).total_seconds()
                        }
                    )
                except Exception as e:
                    logger.log_operation(
                        operation="cell_write_failed",
                        file_path=file_path,
                        success=False,
                        error=e,
                        details={
                            'coordinates': f"Row: {row}, Column: {column}",
                            'attempted_value': value
                        }
                    )
                    raise

            def input_value():
                return '✓ ' + current_time.strftime('%Y/%m/%d,%H:%M:%S')

            # Execute the operations
            row, col = Find_the_cell()
            write_to_cell(file_path, row, col, input_value())

            logger.log_operation(
                operation="complete_operation",
                file_path=file_path,
                details={
                    'total_duration': (datetime.now() - start_time).total_seconds()
                }
            )

        except Exception as e:
            logger.log_operation(
                operation="workbook_operation_failed",
                file_path=file_path,
                success=False,
                error=e
            )
            raise

    except Exception as e:
        logger.log_operation(
            operation="file_processing_failed",
            file_path="unknown",
            success=False,
            error=e,
            details={'input_text': import_text}
        )
        raise


@app.post("/transcribe", response_model=TranscriptionResponse)
async def transcribe_cantonese(request: TranscriptionRequest):
    try:
        # 1. Call Pollinations.ai API directly
        response = requests.post(
            "https://text.pollinations.ai/openai",
            json={
                "messages": [{
                    "role": "user",
                    "content": [
                        {"type": "text",
                         "text": "Transcribe this Cantonese audio to Traditional Chinese,the number should be the arabic numbers"},
                        {
                            "type": "input_audio",
                            "input_audio": {
                                "data": request.audio_data,
                                "format": request.format
                            }
                        }
                    ]
                }],
                "model": "openai-audio",
                "language": "yue"
            },
            timeout=30
        )

        # 2. Validate response
        result = response.json()
        if "choices" not in result:
            raise HTTPException(502, "Invalid API response")

        return {"text": result['choices'][0]['message']['content']}

    except requests.exceptions.RequestException as e:
        raise HTTPException(502, f"API request failed: {str(e)}")
    except Exception as e:
        raise HTTPException(500, f"Processing failed: {str(e)}")


class TextData(BaseModel):
    text: str


@app.post("/confirmation")
async def confirmation(data: TextData):
    print(data)
    print(type(data))
    print(type(str(data)))
    data, num = extract_and_convert_chinese_numbers(str(data))

    def contains_number(s):
        return bool(re.search(r'\d', s))

    if contains_number(data):
        read_and_write(data)
        print('wrote')
    return {"status": "success"}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
